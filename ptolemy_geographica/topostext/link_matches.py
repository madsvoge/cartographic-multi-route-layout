#!/usr/bin/env python3
"""
Cross-reference topostext_209.csv and the annotated catalogue with a
single fuzzy coordinate+name score, and write the result back into both
CSVs. Supersedes the old strict-only (0.02deg) matching: that's still
exactly how an exact match scores (100%), but a citation with no exact
coordinate match now gets a real chance too, scored rather than just
binary yes/no.

Scoring (see score_match())
----------------------------
Distance is Manhattan degrees (lon diff + lat diff), the same metric used
throughout this project (crossref_topostext.py's _MATCH_TOL_DEG, etc).

Two distance thresholds, not one - a single cutoff used for both "is this
even a candidate" and "how much does distance count" was too blunt: it
missed real matches like catalogue "Garra" (4.02.25.04, 15°10'/32°50')
against topostext's own "Garra" (4.2.25.4, 16°30'/32°50') - the *exact
same name*, 1.33deg apart (one source's transcription drifted a degree
and a bit), rejected outright by an old 1.2deg hard cutoff before the
name was even looked at.

- distance <= EXACT_TOL_DEG (0.02): score = 100 outright. Both sources
  round DMS the same way, so anything this close is the same point,
  full stop - no need for the name to agree (topostext's phrase is often
  noisy lead-in prose, not a clean name, on a section's first citation).
- distance > CANDIDATE_WINDOW_DEG (3.0): not considered a candidate at
  all, regardless of name - a coincidentally-identical short name
  somewhere unrelated across a whole book shouldn't out-vote real
  geography entirely.
- EXACT_TOL_DEG < distance <= CANDIDATE_WINDOW_DEG: score blends a
  distance component - linear falloff from just-under-1 to 0 across the
  *tighter* DIST_SCORE_REF_DEG (1.5), so it can bottom out at 0 well
  before the wider candidate window does, rather than being generous all
  the way out to 3deg - and a name-similarity component (see
  _name_similarity in verify_near_matches.py - graduated token
  similarity after translating the catalogue's German descriptor
  vocabulary to English and stripping topostext's own multi-city-run
  lead-in prose, tolerant of a shared stem, plural, or transliteration
  drift like "isca"/"iscas" or "taurische"/"taurianus"), 55/45 weighted
  toward distance since coordinates are the primary evidence. A near-
  identical name can still carry a candidate over the match threshold on
  its own even once the distance component has floored to 0 (Garra: dist
  component ~0.11, name component 1.0 -> score 51). A small bonus is
  added when the phrase's implied type (crossref_topostext.py's
  _TYPE_HINTS - "mouth of"/"estuary" implies river_mouth/coast/harbor,
  etc.) matches the candidate's actual category, to break ties between
  two real, differently-named points that happen to sit close together
  (a plain city entry right next to the river-mouth point a phrase like
  "mouth of the X river" is actually describing).

A candidate below MATCH_THRESHOLD (45) doesn't set *_matched to "yes" -
that would claim more confidence than it deserves - but the score and the
best candidate found are still written to *_match_score/*_name/*_ref
regardless, rather than left blank, so a near-miss is visible instead of
looking identical to "nothing plausible nearby at all". unmapped_review.xlsx
carries the same score for exactly that reason, sorted with the closest
near-misses first.

Usage
-----
    python3 link_matches.py
"""

from __future__ import annotations

import csv
import re
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

from verify_near_matches import _name_similarity

SCRIPT_DIR = Path(__file__).resolve().parent
CATALOGUE = SCRIPT_DIR.parent / "data" / "ptolemy_catalogue_annotated.csv"
TOPOSTEXT = SCRIPT_DIR / "topostext_209.csv"
REVIEW_XLSX = SCRIPT_DIR / "unmapped_review.xlsx"

EXACT_TOL_DEG = 0.02
CANDIDATE_WINDOW_DEG = 3.0  # how far a candidate can be considered at all
DIST_SCORE_REF_DEG = 1.5  # where the distance component of the score bottoms out (tighter than the window above)
MATCH_THRESHOLD = 45.0
_TYPE_BONUS = 8.0

# Same as crossref_topostext.py's _TYPE_HINTS - reused here as a tie-
# breaker between two real, differently-named points sitting close
# together, not as a hard filter.
_TYPE_HINTS: list[tuple[re.Pattern, set[str]]] = [
    (re.compile(r"\bmouth of\b|\bestuary\b", re.IGNORECASE), {"river_mouth", "coast", "harbor"}),
    (re.compile(r"\bisland[s]?\b", re.IGNORECASE), {"island"}),
    (re.compile(r"\bpromontory\b|\bcape\b", re.IGNORECASE), {"coast", "mountain"}),
    (re.compile(r"\bharbor\b|\bharbour\b|\bport\b", re.IGNORECASE), {"harbor", "coast"}),
    (re.compile(r"\bbay\b", re.IGNORECASE), {"coast"}),
    (re.compile(r"\bcity\b|\btown\b", re.IGNORECASE), {"city"}),
    (re.compile(r"\bsource[s]?\b|\bquelle\b", re.IGNORECASE), {"river"}),
    (re.compile(r"\blake\b", re.IGNORECASE), {"lake"}),
    (re.compile(r"\bmountain[s]?\b|\bmount\b", re.IGNORECASE), {"mountain"}),
]

# book -> covered map numbers (see build_labels.py/coverage_summary.py's
# docstrings re: book 4's "map 9" source quirk).
_COVERED_MAPS = {
    "2": set(range(2, 17)),
    "3": set(range(1, 16)),
    "4": set(range(1, 9)),
    "5": set(range(1, 20)),
}


def _expected_categories(phrase: str) -> set[str] | None:
    for pattern, expected in _TYPE_HINTS:
        if pattern.search(phrase):
            return expected
    return None


def score_match(distance: float, phrase: str, name: str, category: str) -> float:
    if distance > CANDIDATE_WINDOW_DEG:
        return 0.0
    if distance <= EXACT_TOL_DEG:
        return 100.0
    dist_component = max(0.0, 1.0 - distance / DIST_SCORE_REF_DEG)
    name_component = _name_similarity(phrase, name)
    score = 100.0 * (0.55 * dist_component + 0.45 * name_component)
    expected = _expected_categories(phrase)
    if expected is not None and category in expected:
        score += _TYPE_BONUS
    return round(min(score, 100.0), 1)


# Columns this script itself adds - stripped back out of a file's existing
# fieldnames before re-adding them once, so re-running this script twice in
# a row (without annotate_dataset.py/parse_topostext.py regenerating the
# file fresh in between) can't keep appending duplicate columns - a real
# bug found in topostext_209.csv, which had accumulated nine copies of
# these from repeated runs during scoring-algorithm tuning.
_CAT_MATCH_COLS = ["topostext_matched", "topostext_ref", "topostext_name", "topostext_match_score"]
_TOPO_MATCH_COLS = ["catalogue_matched", "catalogue_ref_id", "catalogue_name", "catalogue_match_score"]


def load_rows(path: Path, strip_cols: list[str]) -> tuple[list[str], list[dict]]:
    with path.open(newline="", encoding="utf-8") as fh:
        reader = csv.DictReader(fh)
        fieldnames = [f for f in (reader.fieldnames or []) if f not in strip_cols]
        return fieldnames, list(reader)


def write_rows(path: Path, fieldnames: list[str], rows: list[dict]) -> None:
    with path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def topo_ref(row: dict) -> str:
    return f"{row['book']}.{row['map']}.{row['section']}.{row['position']}"


def _best(lon: float, lat: float, phrase: str, candidates: list[tuple[dict, str, str, str]]) -> tuple[dict, float] | tuple[None, None]:
    """candidates: list of (row, lon_key, lat_key, name_key)."""
    best, best_score = None, -1.0
    for cand, lon_key, lat_key, name_key in candidates:
        d = abs(float(cand[lon_key]) - lon) + abs(float(cand[lat_key]) - lat)
        if d > CANDIDATE_WINDOW_DEG:
            continue
        s = score_match(d, phrase, cand[name_key], cand.get("category", ""))
        if s > best_score:
            best, best_score = cand, s
    return (best, best_score) if best is not None else (None, None)


def main() -> int:
    cat_fields, cat_rows = load_rows(CATALOGUE, _CAT_MATCH_COLS)
    topo_fields, topo_rows = load_rows(TOPOSTEXT, _TOPO_MATCH_COLS)

    cat_by_book: dict[str, list[dict]] = {}
    for row in cat_rows:
        if row.get("ref_id") and row.get("lon_ptolemy"):
            cat_by_book.setdefault(row["ref_id"].split(".")[0], []).append(row)

    topo_by_book: dict[str, list[dict]] = {}
    for row in topo_rows:
        topo_by_book.setdefault(row["book"], []).append(row)

    # --- catalogue -> topostext ---
    # Columns always carry the *best candidate found*, even below the match
    # threshold - topostext_matched is the yes/no decision, but a sub-
    # threshold near-miss and its score are still useful to see (in
    # unmapped_review.xlsx especially) rather than showing a blank next to
    # a row that in fact had a plausible, just-not-quite-good-enough lead.
    for row in cat_rows:
        row["topostext_matched"] = "no"
        row["topostext_ref"] = ""
        row["topostext_name"] = ""
        row["topostext_match_score"] = ""
        if not (row.get("ref_id") and row.get("lon_ptolemy")):
            continue
        book = row["ref_id"].split(".")[0]
        lon, lat = float(row["lon_ptolemy"]), float(row["lat_ptolemy"])
        candidates = [(c, "lon_decimal", "lat_decimal", "name_phrase") for c in topo_by_book.get(book, [])]
        best, score = _best(lon, lat, row["name"], candidates)
        if best is not None:
            row["topostext_ref"] = topo_ref(best)
            row["topostext_name"] = best["name_phrase"]
            row["topostext_match_score"] = score
            if score >= MATCH_THRESHOLD:
                row["topostext_matched"] = "yes"

    # --- topostext -> catalogue ---
    for row in topo_rows:
        row["catalogue_matched"] = "no"
        row["catalogue_ref_id"] = ""
        row["catalogue_name"] = ""
        row["catalogue_match_score"] = ""
        lon, lat = float(row["lon_decimal"]), float(row["lat_decimal"])
        candidates = [(c, "lon_ptolemy", "lat_ptolemy", "name") for c in cat_by_book.get(row["book"], [])]
        best, score = _best(lon, lat, row["name_phrase"], candidates)
        if best is not None:
            row["catalogue_ref_id"] = best["ref_id"]
            row["catalogue_name"] = best["name"]
            row["catalogue_match_score"] = score
            if score >= MATCH_THRESHOLD:
                row["catalogue_matched"] = "yes"

    new_cat_fields = cat_fields + _CAT_MATCH_COLS
    new_topo_fields = topo_fields + _TOPO_MATCH_COLS
    write_rows(CATALOGUE, new_cat_fields, cat_rows)
    write_rows(TOPOSTEXT, new_topo_fields, topo_rows)

    cat_matched = sum(1 for r in cat_rows if r["topostext_matched"] == "yes")
    topo_matched = sum(1 for r in topo_rows if r["catalogue_matched"] == "yes")
    print(f"catalogue: {len(cat_rows)} rows, {cat_matched} marked topostext_matched=yes (score >= {MATCH_THRESHOLD}) -> {CATALOGUE.name}")
    print(f"topostext: {len(topo_rows)} rows, {topo_matched} marked catalogue_matched=yes (score >= {MATCH_THRESHOLD}) -> {TOPOSTEXT.name}")

    unmapped_cat = [
        r
        for r in cat_rows
        if r["topostext_matched"] == "no"
        and r.get("ref_id")
        and r["ref_id"].split(".")[0].isdigit()
        and int(r["ref_id"].split(".")[1]) in _COVERED_MAPS.get(r["ref_id"].split(".")[0], set())
    ]
    unmapped_topo = [r for r in topo_rows if r["catalogue_matched"] == "no"]

    # Closest near-misses first - the score column is exactly what makes a
    # row worth a second look (a 40 just missed the cut; a 3 has nothing
    # plausible nearby at all), so sort by it rather than catalogue order.
    unmapped_cat.sort(key=lambda r: float(r["topostext_match_score"] or 0), reverse=True)
    unmapped_topo.sort(key=lambda r: float(r["catalogue_match_score"] or 0), reverse=True)

    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "unmapped_catalogue"
    cat_cols = [
        "ref_id", "name", "category", "book", "tabula", "modern_location", "lon_ptolemy", "lat_ptolemy",
        "topostext_match_score", "topostext_name", "topostext_ref",
    ]
    ws1.append(cat_cols)
    for r in unmapped_cat:
        ws1.append([r.get(c, "") for c in cat_cols])
    for i, c in enumerate(cat_cols, start=1):
        ws1.column_dimensions[get_column_letter(i)].width = max(12, len(c) + 2)

    ws2 = wb.create_sheet("unmapped_topostext")
    topo_cols = [
        "book", "map", "section", "position", "name_phrase", "lon_dms", "lat_dms", "lon_decimal", "lat_decimal",
        "catalogue_match_score", "catalogue_name", "catalogue_ref_id",
    ]
    ws2.append(topo_cols)
    for r in unmapped_topo:
        ws2.append([r.get(c, "") for c in topo_cols])
    for i, c in enumerate(topo_cols, start=1):
        ws2.column_dimensions[get_column_letter(i)].width = max(12, len(c) + 2)

    wb.save(REVIEW_XLSX)
    print(f"{len(unmapped_cat)} unmapped catalogue points (scoped to topostext's covered range) -> sheet 'unmapped_catalogue'")
    print(f"{len(unmapped_topo)} unmapped topostext citations -> sheet 'unmapped_topostext'")
    print(f"-> {REVIEW_XLSX}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
